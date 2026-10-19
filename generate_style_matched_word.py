import os
from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.shared import Inches, Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import qn, nsdecls

# Paths
out_dir = r"C:\Users\rapee\Documents\GitHub\cpe_air_remote"
assets_dir = os.path.join(out_dir, "word_assets")
os.makedirs(assets_dir, exist_ok=True)

img_wired_src = r"C:\Users\rapee\.gemini\antigravity\brain\d65e9246-7513-4315-acfd-6984cf5f79f3\.user_uploaded\media_1792318100591_3eedc695.jpg"
img_ir_src = r"C:\Users\rapee\.gemini\antigravity\brain\d65e9246-7513-4315-acfd-6984cf5f79f3\.user_uploaded\media_1792318114604_a2b2cddc.png"
img_occ_src = r"C:\Users\rapee\.gemini\antigravity\brain\d65e9246-7513-4315-acfd-6984cf5f79f3\.user_uploaded\media_1792318121777_db0131cd.png"

def create_annotated_image(src_path, target_w, target_h, markers, out_name):
    """Crop/resize image to target_w x target_h, add dark frame and numbered orange markers."""
    im = Image.open(src_path).convert("RGBA")
    # Center-crop to aspect ratio
    src_w, src_h = im.size
    target_ratio = target_w / target_h
    src_ratio = src_w / src_h
    if src_ratio > target_ratio:
        new_w = int(src_h * target_ratio)
        offset = (src_w - new_w) // 2
        im = im.crop((offset, 0, offset + new_w, src_h))
    else:
        new_h = int(src_w / target_ratio)
        offset = (src_h - new_h) // 2
        im = im.crop((0, offset, src_w, offset + new_h))
    im = im.resize((target_w * 2, target_h * 2), Image.Resampling.LANCZOS)

    pad = 12
    canvas = Image.new("RGBA", (target_w * 2 + pad * 2, target_h * 2 + pad * 2), (24, 24, 27, 255))
    canvas.paste(im, (pad, pad))
    draw = ImageDraw.Draw(canvas)

    try:
        font = ImageFont.truetype("C:/Windows/Fonts/arialbd.ttf", 26)
    except:
        font = ImageFont.load_default()

    for num_str, x, y in markers:
        cx = x * 2 + 24
        cy = y * 2 + 24
        r = 24
        # White outer border + orange circle
        draw.ellipse((cx - r - 3, cy - r - 3, cx + r + 3, cy + r + 3), fill=(255, 255, 255, 255))
        draw.ellipse((cx - r, cy - r, cx + r, cy + r), fill=(249, 115, 22, 255))
        bbox = draw.textbbox((0, 0), num_str, font=font)
        tw = bbox[2] - bbox[0]
        th = bbox[3] - bbox[1]
        draw.text((cx - tw / 2, cy - th / 2 - 3), num_str, fill=(255, 255, 255, 255), font=font)

    out_path = os.path.join(assets_dir, out_name)
    canvas.convert("RGB").save(out_path, quality=95)
    return out_path

# Create annotated photos matching the PDF markers
path_fig_ir = create_annotated_image(
    img_ir_src, 480, 330,
    [("1", 115, 46), ("2", 295, 76), ("3", 310, 10), ("4", 398, 50),
     ("5", 295, 142), ("6", 418, 142), ("7", 275, 188), ("8", 92, 248)],
    "fig_3_1_ir.png"
)

path_fig_wired = create_annotated_image(
    img_wired_src, 470, 325,
    [("1", 215, 65), ("2", 280, 175), ("3", 195, 155), ("4", 145, 195),
     ("5", 245, 235), ("6", 95, 215), ("7", 305, 145)],
    "fig_4_1_wired.png"
)

path_fig_occ = create_annotated_image(
    img_occ_src, 340, 265,
    [("1", 232, 34), ("2", 232, 145), ("3", 178, 162), ("4", 72, 178),
     ("5", 76, 32), ("6", 232, 238)],
    "fig_5_1_occ.png"
)

# Build the Word Document
doc = Document()

for section in doc.sections:
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(1.8)
    section.bottom_margin = Cm(1.8)
    section.left_margin = Cm(2.0)
    section.right_margin = Cm(2.0)
    section.different_first_page_header_footer = True

    # Footer on normal pages
    footer = section.footer
    p_foot = footer.paragraphs[0]
    p_foot.alignment = WD_ALIGN_PARAGRAPH.LEFT
    r_f = p_foot.add_run("RMUTT Save Energy · คู่มือการประกอบและติดตั้งฮาร์ดแวร์")
    r_f.font.name = "TH Sarabun New"
    r_f.font.size = Pt(12)
    r_f.font.color.rgb = RGBColor(107, 114, 128)

FONT_BODY = "TH Sarabun New"
FONT_HEAD = "TH Sarabun New"

def style_run(run, size_pt=15.5, bold=False, italic=False, color_rgb=(31, 41, 55), font_name=FONT_BODY):
    run.font.name = font_name
    run.font.size = Pt(size_pt)
    run.bold = bold
    run.italic = italic
    run.font.color.rgb = RGBColor(*color_rgb)
    rPr = run._element.get_or_add_rPr()
    rFonts = rPr.find(qn('w:rFonts'))
    if rFonts is None:
        rFonts = OxmlElement('w:rFonts')
        rPr.insert(0, rFonts)
    rFonts.set(qn('w:ascii'), font_name)
    rFonts.set(qn('w:hAnsi'), font_name)
    rFonts.set(qn('w:cs'), font_name)
    rFonts.set(qn('w:eastAsia'), font_name)
    szCs = rPr.find(qn('w:szCs'))
    if szCs is None:
        szCs = OxmlElement('w:szCs')
        rPr.append(szCs)
    szCs.set(qn('w:val'), str(int(size_pt * 2)))

def set_cell_bg(cell, hex_color):
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{hex_color}"/>')
    cell._tc.get_or_add_tcPr().append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def add_chapter_heading(num_str, title_str):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(10)
    p.paragraph_format.space_after = Pt(8)
    p.paragraph_format.keep_with_next = True
    # Green bottom border on paragraph
    pPr = p._p.get_or_add_pPr()
    pBdr = parse_xml(f'<w:pBdr {nsdecls("w")}><w:bottom w:val="single" w:sz="18" w:space="4" w:color="16A34A"/></w:pBdr>')
    pPr.append(pBdr)
    if num_str:
        r_num = p.add_run(f"{num_str} ")
        style_run(r_num, size_pt=22, bold=True, color_rgb=(22, 163, 74))
    r_title = p.add_run(title_str)
    style_run(r_title, size_pt=22, bold=True, color_rgb=(31, 41, 55))
    return p

def add_sec_heading(title_str):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.keep_with_next = True
    r = p.add_run(title_str)
    style_run(r, size_pt=17.5, bold=True, color_rgb=(21, 128, 61))
    return p

def add_sub_heading(title_str):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.keep_with_next = True
    r = p.add_run(title_str)
    style_run(r, size_pt=16, bold=True, color_rgb=(17, 24, 39))
    return p

def add_p(text):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.15
    r = p.add_run(text)
    style_run(r, size_pt=15.5)
    return p

def add_bullet(bold_prefix, rest_text):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Cm(0.6)
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.line_spacing = 1.15
    r0 = p.add_run("• ")
    style_run(r0, size_pt=15.5, bold=True, color_rgb=(21, 128, 61))
    if bold_prefix:
        r1 = p.add_run(bold_prefix)
        style_run(r1, size_pt=15.5, bold=True)
    if rest_text:
        r2 = p.add_run(rest_text)
        style_run(r2, size_pt=15.5)
    return p

def add_step(step_num, bold_prefix, rest_text):
    """Create a numbered step row using a borderless 2-column mini layout or green circle badge."""
    circle_chars = {1: "➊", 2: "➋", 3: "➌", 4: "➍", 5: "➎", 6: "➏"}
    badge = circle_chars.get(step_num, f"({step_num})")
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Cm(0.8)
    p.paragraph_format.first_line_indent = Cm(-0.8)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.15
    r_b = p.add_run(f"{badge}  ")
    style_run(r_b, size_pt=16, bold=True, color_rgb=(21, 128, 61))
    if bold_prefix:
        r1 = p.add_run(bold_prefix)
        style_run(r1, size_pt=15.5, bold=True)
    if rest_text:
        r2 = p.add_run(rest_text)
        style_run(r2, size_pt=15.5)
    return p

def add_clean_table(headers, rows, col_widths_cm=None):
    """Dark header (#27272A) with white bold text and horizontal row dividers (#E5E7EB), matching the PDF."""
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False

    # Set clean horizontal borders only (no vertical gridlines)
    tblPr = table._tbl.tblPr
    borders = parse_xml(f'''
        <w:tblBorders {nsdecls("w")}>
            <w:top w:val="single" w:sz="8" w:space="0" w:color="27272A"/>
            <w:bottom w:val="single" w:sz="8" w:space="0" w:color="CBD5E1"/>
            <w:insideH w:val="single" w:sz="6" w:space="0" w:color="E5E7EB"/>
            <w:left w:val="none"/>
            <w:right w:val="none"/>
            <w:insideV w:val="none"/>
        </w:tblBorders>
    ''')
    tblPr.append(borders)

    # Header Row
    hdr_cells = table.rows[0].cells
    for i, h_text in enumerate(headers):
        set_cell_bg(hdr_cells[i], "27272A")
        set_cell_margins(hdr_cells[i], top=90, bottom=90, left=130, right=130)
        p = hdr_cells[i].paragraphs[0]
        p.paragraph_format.space_before = Pt(1)
        p.paragraph_format.space_after = Pt(1)
        r = p.add_run(h_text)
        style_run(r, size_pt=15, bold=True, color_rgb=(255, 255, 255))

    # Data Rows
    for r_idx, r_data in enumerate(rows):
        row_cells = table.rows[r_idx + 1].cells
        for c_idx, val in enumerate(r_data):
            set_cell_margins(row_cells[c_idx], top=80, bottom=80, left=130, right=130)
            p = row_cells[c_idx].paragraphs[0]
            p.paragraph_format.space_before = Pt(1)
            p.paragraph_format.space_after = Pt(1)
            p.paragraph_format.line_spacing = 1.1
            r = p.add_run(val)
            style_run(r, size_pt=14.5, bold=(c_idx == 0))

    if col_widths_cm:
        for row in table.rows:
            for idx, w_cm in enumerate(col_widths_cm):
                row.cells[idx].width = Cm(w_cm)

    p_sp = doc.add_paragraph()
    p_sp.paragraph_format.space_after = Pt(2)
    return table

def add_callout_box(title_text, body_text, style_type="info"):
    """Single-cell table styled like the callout boxes in the PDF."""
    colors = {
        "info": ("F1F5F9", "0284C7"),
        "warn": ("FFFBEB", "D97706"),
        "danger": ("FEF2F2", "DC2626")
    }
    bg_hex, border_hex = colors.get(style_type, ("F1F5F9", "0284C7"))
    t = doc.add_table(rows=1, cols=1)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = t.cell(0, 0)
    cell.width = Cm(17.0)
    set_cell_bg(cell, bg_hex)
    set_cell_margins(cell, top=120, bottom=120, left=180, right=160)
    tcPr = cell._tc.get_or_add_tcPr()
    tcBorders = parse_xml(f'''
        <w:tcBorders {nsdecls("w")}>
            <w:left w:val="single" w:sz="24" w:space="0" w:color="{border_hex}"/>
            <w:top w:val="none"/>
            <w:right w:val="none"/>
            <w:bottom w:val="none"/>
        </w:tcBorders>
    ''')
    tcPr.append(tcBorders)

    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.line_spacing = 1.15
    r_t = p.add_run(title_text + "\n")
    style_run(r_t, size_pt=15.5, bold=True, color_rgb=(15, 23, 42))
    r_b = p.add_run(body_text)
    style_run(r_b, size_pt=14.5, color_rgb=(31, 41, 55))

    p_sp = doc.add_paragraph()
    p_sp.paragraph_format.space_after = Pt(2)

def add_figure_img(img_path, caption_str, width_cm=12.5):
    p_img = doc.add_paragraph()
    p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_img.paragraph_format.space_before = Pt(4)
    p_img.paragraph_format.space_after = Pt(2)
    run_img = p_img.add_run()
    run_img.add_picture(img_path, width=Cm(width_cm))

    p_cap = doc.add_paragraph()
    p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_cap.paragraph_format.space_after = Pt(6)
    r_c = p_cap.add_run(caption_str)
    style_run(r_c, size_pt=14.5, color_rgb=(75, 85, 99))

# ==============================================================================
# PAGE 1: COVER PAGE (Styled Dark Card Table in Word so it is 100% Editable!)
# ==============================================================================
cover_table = doc.add_table(rows=1, cols=1)
cover_table.alignment = WD_TABLE_ALIGNMENT.CENTER
c_cell = cover_table.cell(0, 0)
c_cell.width = Cm(17.0)
set_cell_bg(c_cell, "272C33")
set_cell_margins(c_cell, top=500, bottom=500, left=400, right=400)

p_c = c_cell.paragraphs[0]
p_c.paragraph_format.space_after = Pt(12)
r_logo = p_c.add_run("R🌿\n")
style_run(r_logo, size_pt=44, bold=True, color_rgb=(255, 255, 255))

r_ctitle = p_c.add_run("คู่มือการประกอบฮาร์ดแวร์\nระบบ RMUTT Save Energy\n")
style_run(r_ctitle, size_pt=32, bold=True, color_rgb=(255, 255, 255))

r_ctag = p_c.add_run("\nS A V E   E N E R G Y\n\n")
style_run(r_ctag, size_pt=16, bold=True, color_rgb=(52, 211, 153))

r_cdesc = p_c.add_run(
    "การเตรียมอุปกรณ์ การบัดกรีวงจร การเดินสายในกล่อง และการอัปโหลดเฟิร์มแวร์ ESP32\n"
    "สำหรับบอร์ดควบคุมแอร์ (รีโมต IR และสายแอร์ Daikin SLM8) และเซ็นเซอร์ตรวจจับคน\n"
    "อาคาร 16 · คณะวิศวกรรมศาสตร์ มหาวิทยาลัยเทคโนโลยีราชมงคลธัญบุรี\n\n\n\n\n\n"
)
style_run(r_cdesc, size_pt=17, color_rgb=(243, 244, 246))

r_cdiv = p_c.add_run("────────────────────────────────────────────────────────\n")
style_run(r_cdiv, size_pt=14, bold=True, color_rgb=(34, 197, 94))

r_cmeta = p_c.add_run(
    "เอกสารฉบับวิศวกรรมและผู้ดูแลระบบ · จัดทำเมื่อวันที่ 18 ตุลาคม 2569\n"
    "อ้างอิงเฟิร์มแวร์บอร์ดแอร์ IR รุ่น 2.5.2 · บอร์ดสายแอร์รุ่น 2.5.5-slm8 · บอร์ดตรวจจับคนรุ่น 1.3.9\n"
    "ครอบคลุมรายการวัสดุอุปกรณ์ ขั้นตอนการประกอบทีละขั้น ผังขาพิน และการตั้งค่าเข้าสู่ระบบ"
)
style_run(r_cmeta, size_pt=14.5, color_rgb=(209, 213, 219))

# ==============================================================================
# PAGE 2: TABLE OF CONTENTS (สารบัญ)
# ==============================================================================
doc.add_page_break()
add_chapter_heading("", "สารบัญ")

toc_entries = [
    ("1 · ภาพรวมชุดฮาร์ดแวร์ในระบบ", "หน้า 3", "หน้าที่ของกล่องทั้ง 3 แบบ · ความเชื่อมโยงในระบบ · ตารางเปรียบเทียบสเปกแต่ละบอร์ด"),
    ("2 · อุปกรณ์เสริม สายไฟ และเครื่องมือช่างที่ต้องเตรียม", "หน้า 4", "สายไฟบ้าน AC 220V · ขั้วต่อแยกสายไฟ Wago · สายสัญญาณ · แผ่นปริ้นท์ไข่ปลา · เครื่องมือช่าง"),
    ("3 · การประกอบบอร์ดควบคุมแอร์แบบอินฟราเรด (IR)", "หน้า 5 – 6", "รายการอุปกรณ์ · ตำแหน่งภายในกล่อง · การบัดกรีวงจร 2N2222 และไฟ LED · การเดินสายไฟ · ตารางขาพิน"),
    ("4 · การประกอบบอร์ดควบคุมแอร์แบบสายแอร์ (Daikin SLM8)", "หน้า 7 – 8", "รายการอุปกรณ์เพิ่มเติม · วงจรไดโอด 1N4148 และตัวแบ่งแรงดัน 100k/200k · การต่อ DHT11 · การต่อสายแอร์บนผนัง"),
    ("5 · การประกอบกล่องเซ็นเซอร์ตรวจจับคน (HLK-LD2410)", "หน้า 9", "รายการอุปกรณ์ · การบัดกรีไฟ LED ฝาหน้า · การติดตั้งเรดาร์คลื่นมิลลิเมตร LD2410 · ตารางขาพิน"),
    ("6 · การติดตั้งโปรแกรมและอัปโหลดเฟิร์มแวร์ลง ESP32", "หน้า 10", "การเช็ควงจรด้วยมัลติมิเตอร์ · การตั้งค่า Arduino IDE และไลบรารี · การเลือก Partition · การอัปโหลดผ่านสาย USB"),
    ("7 · การตั้งค่าบอร์ดผ่านบลูทูธและการติดตั้งหน้างาน", "หน้า 11", "การจองช่องและสร้างโทเคน · การจับคู่บลูทูธในแอปมือถือ · การคล้องห่วง CT และการตรวจหน้าจอ OLED"),
    ("ภาคผนวก ก · ความหมายไฟแสดงสถานะ LED และหน้าจอ OLED", "หน้า 12", "ไฟสถานะ 3 สีบนบอร์ดแอร์ · ไฟสถานะ 2 สีบนเซ็นเซอร์ตรวจจับคน · สัญลักษณ์บนจอ OLED 1.3 นิ้ว"),
    ("ภาคผนวก ข · การแก้ปัญหาฮาร์ดแวร์และการอัปโหลดโค้ดเบื้องต้น", "หน้า 13", "อาการที่พบบ่อย · สาเหตุ · วิธีตรวจสอบและแก้ไขหน้างาน"),
]

for title_l, page_r, sub_txt in toc_entries:
    p_t = doc.add_paragraph()
    p_t.paragraph_format.space_before = Pt(6)
    p_t.paragraph_format.space_after = Pt(1)
    r_tl = p_t.add_run(f"{title_l}   ({page_r})")
    style_run(r_tl, size_pt=16.5, bold=True, color_rgb=(21, 128, 61))

    p_s = doc.add_paragraph()
    p_s.paragraph_format.space_after = Pt(6)
    pPr = p_s._p.get_or_add_pPr()
    pBdr = parse_xml(f'<w:pBdr {nsdecls("w")}><w:bottom w:val="single" w:sz="6" w:space="4" w:color="E5E7EB"/></w:pBdr>')
    pPr.append(pBdr)
    r_s = p_s.add_run(sub_txt)
    style_run(r_s, size_pt=15, color_rgb=(55, 65, 81))

# ==============================================================================
# PAGE 3: บทที่ 1 ภาพรวมชุดฮาร์ดแวร์ในระบบ
# ==============================================================================
doc.add_page_break()
add_chapter_heading("1", "ภาพรวมชุดฮาร์ดแวร์ในระบบ")
add_sec_heading("ฮาร์ดแวร์ทั้ง 3 แบบทำหน้าที่อะไร")
add_p(
    "ระบบ RMUTT Save Energy ใช้บอร์ดไมโครคอนโทรลเลอร์ ESP32 เป็นหัวใจหลักของอุปกรณ์หน้าห้องเรียน แบ่งออกเป็น 3 รูปแบบ "
    "ตามลักษณะของเครื่องปรับอากาศและหน้าที่ในห้องเรียน โดยทุกกล่องใช้บลูทูธ (BLE) สำหรับตั้งค่าครั้งแรก และสื่อสารกับเซิร์ฟเวอร์ผ่านโปรโตคอล MQTT บนเครือข่าย Wi-Fi ของอาคาร"
)
add_bullet("บอร์ดควบคุมแอร์แบบอินฟราเรด (IR): ", "ใช้กับแอร์ทั่วไปที่สั่งงานด้วยรีโมตไร้สาย (เช่น ห้อง 16303 และ 16307) โดยยิงแสงอินฟราเรด 38 kHz ไปที่หน้าแอร์ และใช้มิเตอร์ PZEM-004T วัดค่าไฟเพื่อยืนยันว่าแอร์ทำงานจริง")
add_bullet("บอร์ดควบคุมแอร์แบบสายแอร์ (Daikin SLM8): ", "ใช้แทนรีโมตติดผนังแบบมีสายของแอร์ Daikin โดยรับและส่งเฟรมข้อมูลผ่านสายแอร์โดยตรงทุก 1 วินาที รู้สถานะคอมเพรสเซอร์จริงจากแอร์ และใช้ DHT11 วัดอุณหภูมิห้อง")
add_bullet("กล่องเซ็นเซอร์ตรวจจับคน (HLK-LD2410): ", "ติดตั้งห้องละ 1 กล่อง ใช้เรดาร์คลื่นมิลลิเมตร 24 GHz ตรวจจับทั้งคนที่กำลังเดินและคนที่นั่งเรียนนิ่งๆ เพื่อส่งให้เซิร์ฟเวอร์ตัดสินใจปิดแอร์อัตโนมัติเมื่อห้องว่าง")

add_sec_heading("ตารางเปรียบเทียบสเปกของฮาร์ดแวร์ทั้ง 3 กล่อง")
add_clean_table(
    ["ชนิดกล่องฮาร์ดแวร์", "ไฟล์เฟิร์มแวร์ (รุ่นล่าสุด)", "อุปกรณ์ที่ต่อพ่วงกับ ESP32", "แหล่งจ่ายไฟเลี้ยง"],
    [
        [
            "1. บอร์ดแอร์แบบ IR\n(1 กล่อง ต่อแอร์ 1 เครื่อง)",
            "rmutt_esp32_air.ino\n(รุ่น 2.5.2)",
            "• วงจร 2N2222 + หลอด IR LED\n• มิเตอร์ PZEM-004T + ห่วง CT\n• จอ OLED 1.3\" + ไฟ LED 3 สี",
            "ไฟบ้าน AC 220V ผ่านโมดูล HDR-15-5 (แปลงเป็น 5V DC)"
        ],
        [
            "2. บอร์ดสายแอร์ Daikin\n(1 กล่อง ต่อแอร์ 1 เครื่อง)",
            "rmutt_esp32_air_wired.ino\n(รุ่น 2.5.5-slm8)",
            "• ไดโอด 1N4148 + ตัวแบ่ง 100k/200k\n• เซ็นเซอร์อุณหภูมิ DHT11\n• มิเตอร์ PZEM-004T + จอ OLED + LED",
            "ไฟบ้าน AC 220V ผ่านโมดูล HDR-15-5 (แปลงเป็น 5V DC)"
        ],
        [
            "3. เซ็นเซอร์ตรวจจับคน\n(1 กล่อง ต่อ 1 ห้องเรียน)",
            "rmutt_esp32_occupancy.ino\n(รุ่น 1.3.9)",
            "• เรดาร์ HLK-LD2410 (UART)\n• ไฟแสดงสถานะ LED 2 สี (เขียว/แดง)",
            "อะแดปเตอร์ USB 5V เสียบพอร์ต USB Type-C ของบอร์ด"
        ]
    ],
    col_widths_cm=[3.8, 4.4, 5.0, 3.8]
)

add_callout_box(
    "เหตุผลที่ใช้ระบบวัดไฟแอร์ 1 เฟส (220V) ร่วมกับ PZEM-004T",
    "โมดูลวัดพลังงาน PZEM-004T ออกแบบมาสำหรับระบบไฟฟ้า 1 เฟส (220V) โดยเฉพาะ ซึ่งสอดคล้องกับเครื่องปรับอากาศในห้องเรียนที่ทดสอบ (16303 และ 16307) ทำให้ใช้ PZEM-004T เพียง 1 ชุดต่อเครื่อง ติดตั้งง่าย ปลอดภัย และไม่ต้องดัดแปลงตู้ไฟ 3 เฟสของอาคาร",
    "info"
)

# ==============================================================================
# PAGE 4: บทที่ 2 อุปกรณ์เสริม สายไฟ และเครื่องมือช่างที่ต้องเตรียม
# ==============================================================================
doc.add_page_break()
add_chapter_heading("2", "อุปกรณ์เสริม สายไฟ และเครื่องมือช่างที่ต้องเตรียม")
add_p(
    "ในการประกอบฮาร์ดแวร์ทั้ง 3 กล่องให้แข็งแรง ปลอดภัย และพร้อมติดตั้งใช้งานจริง นอกจากบอร์ดไมโครคอนโทรลเลอร์และโมดูลเซ็นเซอร์แล้ว ต้องเตรียมสายไฟ ขั้วต่อ แผ่นวงจร และเครื่องมือช่างดังตารางต่อไปนี้"
)

add_clean_table(
    ["หมวดอุปกรณ์เสริม", "รายการและขนาดที่ใช้จริง", "ใช้ทำอะไรในขั้นตอนการประกอบ"],
    [
        [
            "1. สายไฟแรงดันสูง (AC 220V)\n(ใช้ในกล่องแอร์ 1 และ 2)",
            "• สายไฟพร้อมปลั๊กตัวผู้ มอก. ชนิด VCT ขนาด 2×0.75 หรือ 2×1.0 ตร.มม. (ยาว 1.5–2 ม.)\n• สายไฟแกนเดี่ยว/สายอ่อนหุ้มฉนวน PVC (THW หรือ VSF) ขนาด 1.0–1.5 ตร.มม. แยกสีน้ำตาล (Line: L) และสีฟ้า (Neutral: N)",
            "• นำไฟบ้าน 220V เข้าสู่ตัวกล่องควบคุม\n• ใช้เดินสายพ่วงจากตัวแยกสายไฟไปยังช่องอินพุต AC ของภาคจ่ายไฟ HDR-15-5 และขั้ววัดแรงดันของมิเตอร์ PZEM-004T"
        ],
        [
            "2. ขั้วต่อแยกสายไฟ AC\n(ใช้ในกล่องแอร์ 1 และ 2)",
            "• ขั้วต่อสายไฟแบบคันโยก (Wago / Lever-Nut Connector) ชนิด 2 เข้า 2 ออก จำนวน 2 ตัว (สีส้ม-เทา ทนกระแส 32A/250V)",
            "• ใช้เป็นจุดพักและแยกสายไฟ 220V ขั้ว L (1 ตัว) และขั้ว N (1 ตัว) ไปจ่ายให้ Power Supply และ PZEM โดยแน่นหนา ไม่ต้องพันเทปดำ และไม่เกิดประกายไฟ"
        ],
        [
            "3. สายไฟแรงดันต่ำและสายสัญญาณ (DC)",
            "• สายไฟอ่อนตัวนำทองแดงขนาด 22–24 AWG หลากสี (แดง, ดำ, เหลือง, เขียว, ขาว, น้ำเงิน)\n• สายแพรจัมเปอร์ (Dupont Wire)\n• สายไฟคู่ แดง-ดำ (ยาว 1.5–3 ม.)",
            "• สายสีแดงและดำใช้เดินไฟเลี้ยงกระแสตรง (+5V, +3.3V และ GND)\n• สายสีอื่นๆ ใช้เดินสัญญาณ UART, I2C และ GPIO ขันเข้าฐานสกรูของ ESP32\n• สายคู่แดง-ดำใช้ต่อหลอด IR LED ออกไปติดหน้าแอร์"
        ],
        [
            "4. แผ่นวงจรและอะไหล่อิเล็กทรอนิกส์ย่อย",
            "• แผ่นปริ้นท์ไข่ปลา (Perfboard / Prototype PCB) ขนาด 4×6 ซม. และ 2×3 ซม.\n• ตัวต้านทานขนาด 1/4W ค่า 220Ω–330Ω และ 100kΩ\n• ขั้วต่อสาย JST XH 2.54 มม. และท่อหดหุ้มฉนวน",
            "• ใช้เป็นฐานบัดกรีวงจรทรานซิสเตอร์ 2N2222 ขับหลอด IR, วงจรไดโอด 1N4148, วงจรแบ่งแรงดันสายแอร์ และตัวต้านทานจำกัดกระแสของไฟ LED\n• ใช้ท่อหดหุ้มจุดบัดกรีป้องกันสายช็อตกัน"
        ],
        [
            "5. กล่องบรรจุและเครื่องมือช่าง",
            "• กล่องพิมพ์ 3 มิติ (3D Print PLA/PETG) พร้อมฝาหน้าและกรอบยึดจอ OLED\n• หัวแร้งบัดกรี + ตะกั่ว, คีมตัด/ปอกสายไฟ, ไขควงปากแบนขนาดเล็ก (2.0 มม.), มัลติมิเตอร์ดิจิทัล และสายข้อมูล USB Type-C",
            "• ใช้ยึดอุปกรณ์ทั้งหมดลงในกล่องให้เป็นระเบียบ\n• ใช้ไขควงปากแบนขันสกรูบนฐานขยายขา ESP32\n• ใช้มัลติมิเตอร์วัดเช็คการลัดวงจรก่อนจ่ายไฟ และใช้สาย USB Type-C อัปโหลดโค้ดจากคอมพิวเตอร์"
        ]
    ],
    col_widths_cm=[4.0, 6.4, 6.6]
)

add_callout_box(
    "กฎความปลอดภัยสำคัญที่สุดก่อนเริ่มประกอบ",
    "1) กล่องบอร์ดแอร์แบบที่ 1 และ 2 มีไฟบ้าน AC 220V เข้าที่ตัวแยกสาย Wago, ช่องด้านล่างของภาคจ่ายไฟ HDR-15-5 และขั้วฝั่งซ้ายของ PZEM-004T ห้ามเสียบปลั๊กไฟ 220V ขณะกำลังขันสายไฟหรือบัดกรีเด็ดขาด\n"
    "2) ทุกครั้งที่จะเสียบสาย USB Type-C จากคอมพิวเตอร์เพื่ออัปโหลดโค้ดหรือดู Serial Monitor ต้องถอดปลั๊กไฟบ้าน 220V ออกก่อนเสมอ",
    "danger"
)

# ==============================================================================
# PAGE 5: บทที่ 3 การประกอบบอร์ดควบคุมแอร์แบบอินฟราเรด (IR)
# ==============================================================================
doc.add_page_break()
add_chapter_heading("3", "การประกอบบอร์ดควบคุมแอร์แบบอินฟราเรด (IR)")
add_sec_heading("3.1 รายการอุปกรณ์และตำแหน่งภายในกล่องบอร์ดแอร์ IR")
add_p("บอร์ดแอร์ IR ใช้สำหรับแอร์ทั่วไปที่สั่งงานด้วยรีโมตอินฟราเรด หนึ่งกล่องใช้ควบคุมแอร์ 1 เครื่อง โดยมีตำแหน่งการจัดวางอุปกรณ์ภายในกล่องดังนี้")

add_figure_img(path_fig_ir, "รูปที่ 3-1 ตำแหน่งอุปกรณ์และการเดินสายภายในกล่องบอร์ดควบคุมแอร์แบบอินฟราเรด (IR)", width_cm=11.8)

add_clean_table(
    ["#", "ส่วนประกอบในรูปที่ 3-1", "หน้าที่และการติดตั้ง"],
    [
        ["1", "จอ OLED 1.3 นิ้ว (I2C SH1106)", "ขันสกรูยึดกับกรอบด้านในฝาหน้า แสดงสถานะแอร์ อุณหภูมิ กำลังไฟ (W) และสถานะ Wi-Fi"],
        ["2", "ESP32 + ฐานขยายขาขันสกรู", "บอร์ดประมวลผลหลัก เสียบลงบนฐาน ESP32 Terminal Adapter เพื่อขันสกรูสายไฟให้แน่นหนา"],
        ["3", "ไฟสถานะ LED 3 สี (แดง/เหลือง/เขียว)", "ยื่นหัวหลอดออกด้านบนกล่อง แสดงสถานะแอร์ทำงาน (เขียว), นอกคาบ (แดง), และ Wi-Fi (เหลือง)"],
        ["4", "แผงวงจรย่อย (ทรานซิสเตอร์ 2N2222)", "แผ่นปริ้นท์ไข่ปลาขนาด 4×6 ซม. บัดกรีทรานซิสเตอร์ 2N2222 + ตัวต้านทาน 220Ω สำหรับขยายกระแสขับหลอด IR LED"],
        ["5", "มิเตอร์วัดไฟ PZEM-004T v3.0", "วัดแรงดัน (V), กระแส (A), กำลังไฟ (W) และหน่วยไฟสะสม (kWh) ส่งเข้า ESP32 ทาง UART"],
        ["6", "ภาคจ่ายไฟ MeanWell HDR-15-5", "แปลงไฟบ้าน AC 220V เป็นไฟเลี้ยงกระแสตรง DC 5V (2.4A) จ่ายให้บอร์ด ESP32 และโมดูลทั้งหมด"],
        ["7", "ขั้วต่อแยกสายไฟ Wago (สีส้ม)", "รับไฟ 220V (สายสีน้ำตาล L และสีฟ้า N) จากสายปลั๊กหลัก แล้วแยกจ่ายไปให้ HDR-15-5 และ PZEM"],
        ["8", "ห่วงคล้อง CT Sensor + สายหลอด IR", "ห่วง CT สีดำนำไปคล้องสายเมนไฟแอร์ 1 เส้น ส่วนสายคู่แดง-ดำต่อหลอด IR LED ไปติดหน้าแอร์"]
    ],
    col_widths_cm=[1.2, 5.6, 10.2]
)

# ==============================================================================
# PAGE 6: บทที่ 3 (ต่อ) ขั้นตอนการประกอบบอร์ดแอร์ IR ทีละขั้น
# ==============================================================================
doc.add_page_break()
add_sec_heading("3.2 ขั้นตอนการประกอบบอร์ดแอร์ IR ทีละขั้น (เริ่มจากศูนย์จนครบกล่อง)")

add_sub_heading("ขั้นที่ 1 · บัดกรีแผงวงจรย่อยบนแผ่นปริ้นท์ไข่ปลา (ทำนอกกล่องก่อน)")
add_step(1, "บัดกรีหลอดไฟสถานะ LED 3 สี: ", "เสียบหลอด LED สีแดง เหลือง เขียว เรียงที่ขอบบนของแผ่นปริ้นท์ไข่ปลา ดัด ขาสั้น (ขาลบ) ของทั้ง 3 หลอดเชื่อมรวมกันเป็นสาย GND เส้นเดียว ส่วน ขายาว (ขาบวก) ของแต่ละหลอดให้บัดกรีต่ออนุกรมกับตัวต้านทาน 220Ω–330Ω แยกสายออกเป็น 3 เส้น")
add_step(2, "บัดกรีวงจรทรานซิสเตอร์ 2N2222 ขับหลอด IR: ", "หันด้านหน้าตัดแบนของ 2N2222 เข้าหาตัว ขาเรียงจากซ้ายไปขวาคือ ขา 1 (ซ้าย = Emitter) ต่อลง GND · ขา 2 (กลาง = Base) ต่อผ่านตัวต้านทาน 220Ω ไปรอเข้า GPIO 17 · ขา 3 (ขวา = Collector) ต่อกับขั้วลบของหลอด IR LED (940nm) ส่วนขั้วบวกของหลอด IR LED ต่อเข้าไฟเลี้ยง +5V")

add_callout_box(
    "ผังการต่อขา Transistor 2N2222 ขับหลอด IR LED บนแผ่นปริ้นท์ไข่ปลา",
    "[ GPIO 17 ] ──► [ ตัวต้านทาน 220Ω ] ──► ขา 2 (Base) ของ 2N2222\n"
    "• ขา 1 (Emitter - ซ้าย) ──► ต่อลง GND ร่วมกับบอร์ด ESP32\n"
    "• ขา 3 (Collector - ขวา) ──► ต่อเข้า ขั้วลบของหลอด IR LED (ส่วนขั้วบวกของหลอด IR LED ต่อไฟ +5V)",
    "info"
)

add_sub_heading("ขั้นที่ 2 · ยึดอุปกรณ์ลงกล่องและเดินสายไฟแรงดันต่ำ + สายไฟบ้าน 220V")
add_step(3, "ยึดโมดูลลงกล่อง 3D Print: ", "กดบอร์ด ESP32 ลงบนฐานขยายขาขันสกรูแล้วยึดไว้มุมซ้ายบน · ยึดแผ่นปริ้นท์ไข่ปลามุมขวาบน (ให้หัวไฟ LED โผล่พ้นรูด้านบน) · วาง HDR-15-5 ด้านขวาล่าง · วาง PZEM-004T ตรงกลางซ้าย · วางขั้ว Wago มุมซ้ายล่าง · ขันจอ OLED ติดด้านในฝาหน้า")
add_step(4, "ขันสกรูสายไฟแรงดันต่ำเข้าฐาน ESP32: ", "ต่อไฟเลี้ยงจาก HDR-15-5 ช่อง +V เข้า 5V/VIN และช่อง -V เข้า GND จากนั้นต่อสาย PZEM, จอ OLED และแผ่นปริ้นท์ไข่ปลาตามตารางขาพินด้านล่าง")
add_step(5, "เข้าสายไฟบ้าน AC 220V และห่วงคล้อง CT: ", "สอดสายปลั๊ก AC 220V เข้าด้านล่างกล่อง แยกสายสีน้ำตาล (L) เข้า Wago ตัวที่ 1 และสายสีฟ้า (N) เข้า Wago ตัวที่ 2 แล้วพ่วงไฟ L / N ไปเข้าช่อง L / N ด้านล่างของ HDR-15-5 และขั้วที่ 1–2 ของ PZEM-004T ส่วนสายห่วงคล้อง CT ต่อเข้าขั้วที่ 3–4 ของ PZEM-004T")

add_sec_heading("3.3 ตารางขาพินของบอร์ดแอร์ IR (อ้างอิงโค้ด rmutt_esp32_air.ino รุ่น 2.5.2)")
add_clean_table(
    ["ขาพิน ESP32", "ชื่อในโค้ด", "ต่อเข้ากับ", "หมายเหตุ"],
    [
        ["5V (VIN), GND", "—", "HDR-15-5 ช่อง +V และ -V", "ไฟเลี้ยงหลัก 5V DC และกราวด์ร่วมของทั้งวงจร"],
        ["GPIO 17", "IR_LED_PIN", "ขา Base ของ 2N2222 (ผ่าน R 220Ω)", "ส่งสัญญาณพัลส์ความถี่ 38 kHz ขับหลอดส่งแสง IR LED"],
        ["GPIO 16 / 18", "PZEM_RX / TX", "16 ➔ TX ของ PZEM · 18 ➔ RX ของ PZEM", "ต่อไขว้ (ESP32 อ่านค่าไฟจาก PZEM ทุก 2 วินาที)"],
        ["GPIO 21 / 22", "kSda / kScl", "ขา SDA / SCL ของจอ OLED 1.3\"", "จอใช้ไฟเลี้ยงจากขา 3V3 และ GND"],
        ["GPIO 25 / 26 / 27", "GREEN / RED / YELLOW", "LED เขียว (25), แดง (26), เหลือง (27)", "ต่อผ่านตัวต้านทาน 220Ω–330Ω ลง GND ทุกหลอด"],
        ["GPIO 0", "SETUP_BUTTON_PIN", "ปุ่ม BOOT บนบอร์ด ESP32", "กดค้างตอนรีเซ็ตเพื่อบังคับเปิดโหมดตั้งค่าบลูทูธ"]
    ],
    col_widths_cm=[3.2, 4.0, 4.8, 5.0]
)

# ==============================================================================
# PAGE 7: บทที่ 4 การประกอบบอร์ดควบคุมแอร์แบบสายแอร์ (Daikin SLM8)
# ==============================================================================
doc.add_page_break()
add_chapter_heading("4", "การประกอบบอร์ดควบคุมแอร์แบบสายแอร์ (Daikin SLM8)")
add_sec_heading("4.1 บอร์ดสายแอร์ต่างจากบอร์ด IR อย่างไร")
add_p(
    "แอร์บางห้องใช้ รีโมตติดผนังแบบมีสาย (รุ่น Daikin SLM8) บอร์ดรุ่นนี้จึงถูกออกแบบมาให้ติดตั้ง แทนที่รีโมตติดผนังเดิม "
    "โดยคุยกับแอร์ผ่านสายสัญญาณเดิมโดยตรง (ส่งเฟรมคำสั่งทุก 1 วินาที และรับเฟรมสถานะกลับมาจากแอร์ ทำให้รู้ว่าคอมเพรสเซอร์ทำงานจริงหรือไม่) "
    "พร้อมใช้เซ็นเซอร์ DHT11 วัดอุณหภูมิห้องส่งให้แอร์แทนเทอร์มิสเตอร์ในรีโมตผนังเดิม"
)

add_figure_img(path_fig_wired, "รูปที่ 4-1 กล่องบอร์ดควบคุมแอร์แบบสายแอร์ (Daikin SLM8) และการเดินสายขึ้นไปยังแป้นรีโมตติดผนังเดิม", width_cm=11.5)

add_clean_table(
    ["#", "ส่วนประกอบในรูปที่ 4-1", "หน้าที่และการเชื่อมต่อ"],
    [
        ["1", "สายสัญญาณแอร์สีขาว (ไปแป้นผนัง)", "เดินสายจากขั้วต่อ JST ในกล่องขึ้นไปต่อกับสายแอร์เดิมบนผนัง 3 เส้น (สายสั่งงาน TX, สายสถานะ RX, และ GND)"],
        ["2", "แผงวงจรสายแอร์ SLM8 (ปริ้นท์ไข่ปลา)", "บัดกรีไดโอด 1N4148 (สำหรับส่งคำสั่ง) และตัวต้านทาน 100kΩ จำนวน 6 ตัว (ทำวงจรแบ่งแรงดัน 100k/200kΩ 2 ชุด)"],
        ["3", "ESP32 + ฐานขยายขาขันสกรู", "รันเฟิร์มแวร์ rmutt_esp32_air_wired.ino (รุ่น 2.5.5-slm8) แยก FreeRTOS Task ส่งเฟรมสายแอร์ทุก 1 วินาที"],
        ["4", "มิเตอร์วัดไฟ PZEM-004T v3.0", "วัดกำลังไฟแอร์ 1 เฟสเหมือนรุ่น IR (แต่รุ่นสายแอร์ย้ายไปต่อที่ขา GPIO 18 และ GPIO 14)"],
        ["5", "ภาคจ่ายไฟ MeanWell HDR-15-5", "จ่ายไฟเลี้ยง DC 5V ให้บอร์ด (จึงไม่ต้องดึงไฟ +5V จากสายแอร์มาใช้)"],
        ["6", "ขั้วต่อแยกสายไฟ Wago (สีส้ม)", "แยกไฟบ้าน AC 220V จ่ายให้ HDR-15-5 และขั้ววัดแรงดันของ PZEM-004T"],
        ["7", "ไฟสถานะ LED + เซ็นเซอร์ DHT11", "แสดงสถานะหน้ากล่อง และใช้ DHT11 (ขา GPIO 4) วัดอุณหภูมิห้องส่งให้แอร์ใช้ตัดคอมเพรสเซอร์"]
    ],
    col_widths_cm=[1.2, 5.6, 10.2]
)

# ==============================================================================
# PAGE 8: บทที่ 4 (ต่อ) วงจรสายแอร์ SLM8 และตารางขาพิน
# ==============================================================================
doc.add_page_break()
add_sec_heading("4.2 ขั้นตอนการบัดกรีวงจรสายแอร์และการต่อสายเข้ากับแอร์ Daikin")
add_sub_heading("ขั้นที่ 1 · บัดกรีวงจรไดโอด 1N4148 และตัวแบ่งแรงดัน 100kΩ / 200kΩ บนแผ่นปริ้นท์ไข่ปลา")
add_p("สายสัญญาณของแอร์ Daikin มีแรงดันค้างอยู่ราว 4.8–5.0V และมีตัวดึงขึ้น (Pull-up) อ่อนมาก จึงห้ามต่อตรงเข้าขา 3.3V ของ ESP32 ให้บัดกรีวงจรคั่นกลางบนแผ่นปริ้นท์ไข่ปลาตามผังด้านล่าง")

add_callout_box(
    "ผังวงจรเชื่อมต่อสายแอร์ Daikin SLM8 บนแผ่นปริ้นท์ไข่ปลา (ใช้ไดโอด 1N4148 + ตัวต้านทาน 100kΩ 6 ตัว)",
    "1) ขาส่งคำสั่ง (GPIO 33): [ GPIO 33 ] ◄──▎[ ไดโอด 1N4148 หันแถบดำเข้า GPIO 33 ] ──► สายสั่งงานแอร์ (สีขาว TX)\n"
    "2) ขารับสถานะ (GPIO 35): สายสถานะแอร์ (สีเหลือง TX/RCV) ──► [ R 100kΩ ] ──┬──► เข้าขา [ GPIO 35 ]\n"
    "                                                                        └──► [ R 200kΩ (ใช้ 100kΩ 2 ตัวอนุกรมกัน) ] ──► ลง GND\n"
    "3) ขาตรวจสัญญาณตัวเอง (GPIO 34): ต่อวงจรแบ่งแรงดัน 100kΩ / 200kΩ ลง GND เหมือน GPIO 35 อีก 1 ชุด แต่ดึงมาจาก สายสั่งงานแอร์ (TX)",
    "info"
)

add_step(1, "ต่อเซ็นเซอร์วัดอุณหภูมิห้อง DHT11 (โมดูล KY-015): ", "ขันสกรูต่อขา DATA เข้าที่ GPIO 4, ขา + (VCC) เข้า 3V3 และขา - (GND) เข้า GND ของฐาน ESP32")
add_step(2, "ต่อสายมิเตอร์ PZEM-004T และไฟ LED (ระวังขาต่างจากรุ่น IR!): ", "ในบอร์ดสายแอร์ ให้ต่อขา TX ของ PZEM เข้าที่ GPIO 18, ขา RX ของ PZEM เข้าที่ GPIO 14 และสายไฟ LED สีเหลืองเข้าที่ GPIO 23 (ส่วนไฟเขียว=25, แดง=26 และจอ OLED=21/22 เหมือนเดิม)")
add_step(3, "การต่อเข้ากับสายแอร์บนผนังห้องเรียน: ", "ถอดรีโมตติดผนัง Daikin SLM8 เดิมออก แล้วต่อสายจากกล่องของเราเข้ากับสายแอร์เพียง 3 เส้น คือ 1) สายสั่งงาน (ขาว TX), 2) สายสถานะ (เหลือง TX/RCV) และ 3) สายกราวด์ (ดำ GND — ต้องต่อร่วมกับ GND ของ ESP32) ส่วนสายไฟ +5V (สีแดง) ของแอร์ ห้ามต่อเด็ดขาด ให้พันเทปเก็บไว้")

add_sec_heading("4.3 ตารางขาพินของบอร์ดสายแอร์ (อ้างอิงโค้ด rmutt_esp32_air_wired.ino รุ่น 2.5.5-slm8)")
add_clean_table(
    ["ขาพิน ESP32", "ชื่อในโค้ด", "ต่อเข้ากับ", "ความแตกต่างจากรุ่น IR"],
    [
        ["GPIO 33", "BUS_TX_PIN", "ไดโอด 1N4148 (หันแถบดำเข้า GPIO 33) ➔ สายสั่งงานแอร์", "ดึงสายลง LOW ส่งเฟรมคำสั่ง 8 ไบต์ทุก 1 วินาที"],
        ["GPIO 35", "BUS_RX_PIN", "สายสถานะแอร์ ผ่านตัวแบ่งแรงดัน 100kΩ / 200kΩ ลง GND", "รับเฟรมสถานะจากแอร์ (รู้ว่าสายตอบและคอมเพรสเซอร์ติดหรือไม่)"],
        ["GPIO 34", "BUS_ECHO_PIN", "สายสั่งงานแอร์ ผ่านตัวแบ่งแรงดัน 100kΩ / 200kΩ ลง GND", "อ่านสัญญาณที่ตัวเองส่งออกไปเพื่อตรวจสอบความถูกต้อง"],
        ["GPIO 4", "DHT_PIN", "ขา DATA ของโมดูล DHT11 (ไฟเลี้ยง 3V3 / GND)", "วัดอุณหภูมิห้องส่งให้แอร์แทนเทอร์มิสเตอร์ในรีโมตผนัง"],
        ["GPIO 18 / 14", "PZEM_RX / TX", "18 ➔ TX ของ PZEM · 14 ➔ RX ของ PZEM", "ย้ายขามาที่ 18 และ 14 (รุ่น IR ใช้ 16 และ 18)"],
        ["GPIO 25 / 26 / 23", "GREEN / RED / YELLOW", "LED เขียว (25), แดง (26), เหลือง (23)", "ไฟเหลืองย้ายมาที่ GPIO 23 (รุ่น IR ใช้ 27)"],
        ["GPIO 21 / 22", "kSda / kScl", "จอ OLED 1.3\" SH1106 (SDA=21, SCL=22)", "แสดงป้าย BUS และ COMP ON (หมุนจอ 180° ให้เอง)"]
    ],
    col_widths_cm=[3.2, 3.8, 5.0, 5.0]
)

# ==============================================================================
# PAGE 9: บทที่ 5 การประกอบกล่องเซ็นเซอร์ตรวจจับคน (HLK-LD2410)
# ==============================================================================
doc.add_page_break()
add_chapter_heading("5", "การประกอบกล่องเซ็นเซอร์ตรวจจับคน (HLK-LD2410)")
add_sec_heading("5.1 ส่วนประกอบภายในกล่องเซ็นเซอร์ตรวจจับคน")
add_p(
    "กล่องเซ็นเซอร์ตรวจจับคนติดตั้งห้องละ 1 กล่อง ใช้เรดาร์คลื่นมิลลิเมตร HLK-LD2410 (24 GHz) อ่านค่าพลังงานนิ่ง (Static Energy) "
    "รายระยะ (Gate 0–8) จึงตรวจจับคนนั่งเรียนนิ่งๆ ได้แม่นยำ และใช้ไฟเลี้ยง 5V จากสาย USB Type-C โดยไม่มีไฟบ้าน 220V ในกล่อง"
)

add_figure_img(path_fig_occ, "รูปที่ 5-1 การจัดวางอุปกรณ์และการเดินสายภายในกล่องเซ็นเซอร์ตรวจจับคน (HLK-LD2410)", width_cm=8.5)

add_sec_heading("5.2 ขั้นตอนการประกอบและเดินสายทีละขั้น")
add_step(1, "บัดกรีแผงไฟสถานะ LED 2 สี ที่ฝาหน้า (หมายเลข 4): ", "ตัดแผ่นปริ้นท์ไข่ปลาขนาด 2×3 ซม. เสียบหลอด LED สีเขียว และ LED สีแดง ให้หัวหลอดโผล่ออกช่องฝาหน้า บัดกรีขาลบ (ขาสั้น) รวมกันเป็นสายสีน้ำตาลต่อเข้า GND ส่วนขาบวก (ขายาว) ต่อผ่านตัวต้านทาน 220Ω–330Ω แยกเป็นสายสีเขียว (เข้า GPIO 25) และสายสีแดง (เข้า GPIO 26)")
add_step(2, "ต่อสายโมดูลเรดาร์ HLK-LD2410 (หมายเลข 1): ", "ที่ขอบบนของบอร์ดเรดาร์สีน้ำเงินมี 5 รูพิน เรียงจากซ้ายไปขวาคือ TX, RX, OUT, GND, VCC ให้ต่อสายเพียง 4 เส้น (เว้นขา OUT ไว้ไม่ต้องต่อ) แล้ววางลงในช่องล็อกด้านบนสุดของกล่องโดยหันด้านสายอากาศสีทองออกด้านหน้า (ตรงกับหน้าต่างฝาหน้าหมายเลข 5)")
add_step(3, "ยึดบอร์ด ESP32 และขันสกรูสายไฟ (หมายเลข 2, 3, 6): ", "เสียบบอร์ด ESP32 ลงบนฐานขยายขาขันสกรู ยึดลงกล่องโดยหันพอร์ต USB Type-C ออกช่องด้านล่างสุด แล้วขันสกรูสายไฟตามตารางด้านล่าง")

add_clean_table(
    ["ขาพิน ESP32", "ชื่อในโค้ด", "ต่อเข้ากับ", "สีสายในรูปที่ 5-1 และหน้าที่"],
    [
        ["5V (VIN)", "—", "ขา VCC ของเรดาร์ HLK-LD2410", "สายสีแดง (จ่ายไฟ 5V จากพอร์ต USB ให้เรดาร์)"],
        ["GND", "—", "ขา GND เรดาร์ + ขาลบร่วมไฟ LED", "สายสีดำ (จากเรดาร์) และสายสีน้ำตาล (จากไฟ LED)"],
        ["GPIO 32", "RADAR_RX_PIN", "ขา TX ของเรดาร์ HLK-LD2410", "ต่อไขว้ · รับข้อมูลเรดาร์ที่ความเร็ว 256,000 baud"],
        ["GPIO 33", "RADAR_TX_PIN", "ขา RX ของเรดาร์ HLK-LD2410", "ต่อไขว้ · ส่งคำสั่งเปิดโหมด Engineering รายเกต"],
        ["GPIO 25 / 26", "GREEN / RED_LED", "LED เขียว (25) และ LED แดง (26)", "เขียว = สถานะเชื่อมต่อเน็ต · แดง = มีคนอยู่ในห้อง"]
    ],
    col_widths_cm=[3.2, 3.8, 4.8, 5.2]
)

# ==============================================================================
# PAGE 10: บทที่ 6 การติดตั้งโปรแกรมและอัปโหลดเฟิร์มแวร์ลง ESP32
# ==============================================================================
doc.add_page_break()
add_chapter_heading("6", "การติดตั้งโปรแกรมและอัปโหลดเฟิร์มแวร์ลง ESP32")
add_sec_heading("6.1 การตรวจเช็ควงจรด้วยมัลติมิเตอร์ก่อนเสียบสายไฟ")
add_p("ก่อนเสียบสาย USB เข้าคอมพิวเตอร์หรือเสียบปลั๊กไฟ 220V ให้ปรับมัลติมิเตอร์ไปที่โหมดวัดความต่อเนื่อง (เสียงร้องเตือน) แล้วตรวจเช็ค 2 จุดเสมอ:")
add_bullet("เช็คฝั่งไฟเลี้ยง DC: ", "วัดระหว่างขั้ว 5V กับ GND และระหว่าง 3V3 กับ GND บนฐาน ESP32 ต้องไม่มีเสียงร้องเตือนลัดวงจร")
add_bullet("เช็คฝั่งไฟบ้าน AC 220V (เฉพาะกล่องแอร์ 1 และ 2): ", "วัดระหว่างขาปลั๊กไฟ L กับ N ต้องไม่ช็อตถึงกัน และต้องไม่รั่วข้ามมายังขั้ว GND ของบอร์ด ESP32")

add_sec_heading("6.2 การเตรียม Arduino IDE และติดตั้งไลบรารี")
add_p("เปิดโปรแกรม Arduino IDE (รุ่น 2.x) ติดตั้งบอร์ดแพ็กเกจ esp32 by Espressif Systems (รุ่น 3.x) ในเมนู Boards Manager จากนั้นไปที่เมนู Library Manager เพื่อติดตั้งไลบรารีให้ครบตามชนิดกล่องดังตารางด้านล่าง")

add_clean_table(
    ["ชนิดกล่องที่จะอัปโหลด", "ไฟล์สเก็ตช์ (.ino) ที่ต้องเปิด", "ไลบรารีที่ต้องติดตั้งใน Arduino IDE"],
    [
        [
            "1. บอร์ดแอร์แบบ IR",
            "firmware/rmutt_esp32_air/\nrmutt_esp32_air.ino (รุ่น 2.5.2)",
            "• ArduinoJson (ต้องใช้รุ่น 6.x เท่านั้น เช่น 6.21.5)\n• PubSubClient · PZEM004Tv30 · U8g2\n• IRremoteESP8266 (รุ่น 2.9.0)"
        ],
        [
            "2. บอร์ดสายแอร์ Daikin",
            "firmware/rmutt_esp32_air_wired/\nrmutt_esp32_air_wired.ino (รุ่น 2.5.5-slm8)",
            "• ArduinoJson (6.x) · PubSubClient · PZEM004Tv30 · U8g2\n• DHT sensor library (by Adafruit) + Adafruit Unified Sensor"
        ],
        [
            "3. เซ็นเซอร์ตรวจจับคน",
            "firmware/rmutt_esp32_occupancy/\nrmutt_esp32_occupancy.ino (รุ่น 1.3.9)",
            "• ArduinoJson (6.x) · PubSubClient\n• ld2410 (by Nick Reynolds — ต้องเป็นรุ่น 0.2.2 ขึ้นไป)"
        ]
    ],
    col_widths_cm=[4.0, 5.5, 7.5]
)

add_callout_box(
    "ห้ามลบไฟล์ประกอบในโฟลเดอร์สเก็ตช์",
    "ในโฟลเดอร์ของแต่ละสเก็ตช์มีไฟล์สำคัญวางอยู่คู่กับไฟล์ .ino ได้แก่ build_opt.h (ช่วยลดขนาดโค้ดให้พอดีแฟลช), partitions.csv, letsencrypt_ca.h, oled_display.h, SaijoDenkiAC.h/.cpp และ DaikinSLM8.h ต้องเก็บไว้ในโฟลเดอร์เดิมครบทุกไฟล์",
    "warn"
)

add_sec_heading("6.3 ขั้นตอนการตั้งค่าและกด Upload โค้ดลงบอร์ด ESP32")
add_step(1, "", "เปิดไฟล์สเก็ตช์ .ino ของกล่องที่ต้องการใน Arduino IDE")
add_step(2, "", "ไปที่เมนู Tools → Board → esp32 เลือกเป็น ESP32 Dev Module")
add_step(3, "", "ไปที่เมนู Tools → Partition Scheme เลือกเป็น No FS 4MB (2MB APP x2) (เพื่อให้โปรแกรมจัดแบ่งพื้นที่แฟลชตามไฟล์ partitions.csv สำหรับรองรับการอัปเดต OTA และเก็บข้อมูลย้อนหลังตอนเน็ตหลุด)")
add_step(4, "ถอดปลั๊กไฟบ้าน 220V ออกก่อนเสมอ ", "แล้วเสียบสาย USB Type-C จากคอมพิวเตอร์เข้าที่บอร์ด ESP32 เลือกพอร์ตในเมนู Tools → Port (เช่น COM3) แล้วกดปุ่ม Upload (ลูกศรชี้ขวา) รอจนขึ้นข้อความ Done uploading.")
add_step(5, "", "กดเปิด Serial Monitor มุมขวาบน ตั้งค่าความเร็วเป็น 115200 baud จะเห็นข้อความบูตระบบและชื่อบลูทูธของบอร์ด เช่น RMUTT-AC-XXXXXXXXXXXX หรือ RMUTT-OCC-XXXXXXXXXXXX แปลว่าบอร์ดพร้อมนำไปตั้งค่าใช้งานแล้ว")

# ==============================================================================
# PAGE 11: บทที่ 7 การตั้งค่าบอร์ดผ่านบลูทูธและการติดตั้งหน้างาน
# ==============================================================================
doc.add_page_break()
add_chapter_heading("7", "การตั้งค่าบอร์ดผ่านบลูทูธและการติดตั้งหน้างาน")
add_sec_heading("7.1 ขั้นตอนการจับคู่บอร์ดเข้ากับห้องเรียนผ่านบลูทูธ (BLE Provisioning)")
add_p("เมื่ออัปโหลดโค้ดเสร็จและนำกล่องไปติดตั้งจ่ายไฟที่ห้องเรียนแล้ว ให้ใช้แอปมือถือ RMUTT Save Energy (เข้าสู่ระบบด้วยบัญชีผู้ดูแลระบบ) เพื่อส่งค่า Wi-Fi และผูกบอร์ดเข้ากับห้องเรียนตามขั้นตอนดังนี้")

add_step(1, "", "ในแอปมือถือ ไปที่แท็บ ควบคุม เลือกห้องเรียนที่ติดตั้งกล่อง แล้วกดปุ่ม เพิ่มอุปกรณ์ ESP32 (สำหรับบอร์ดแอร์ทั้งรุ่น IR และสายแอร์) หรือกด เพิ่มเซ็นเซอร์ (สำหรับกล่องตรวจจับคน) — หรือจองช่องสร้างโทเคน 15 นาทีจากหน้าเว็บแอปก่อนก็ได้")
add_step(2, "", "เปิดบลูทูธบนมือถือแล้วกด ค้นหาบลูทูธ เลือกชื่อบอร์ดที่ตรวจพบ (บอร์ดแอร์ขึ้นต้นด้วย RMUTT-AC-... ส่วนเซ็นเซอร์ตรวจจับคนขึ้นต้นด้วย RMUTT-OCC-...)")
add_step(3, "", "เลือกชื่อเครือข่าย Wi-Fi (2.4 GHz) ที่บอร์ดสแกนพบ แล้วกรอกรหัสผ่าน Wi-Fi")
add_step(4, "", "เลือกหมายเลขแอร์ในห้อง ตั้งชื่อ และเลือก ยี่ห้อ/โปรโตคอลของแอร์ (สำหรับบอร์ดแอร์ IR ให้เลือกให้ตรงกับรีโมตจริง เช่น Daikin, Mitsubishi, Haier, Saijo Denki ส่วน บอร์ดสายแอร์จะถูกล็อกเป็น \"Daikin · สายแอร์\" ให้อัตโนมัติ)")
add_step(5, "", "กดปุ่ม ตั้งค่าและเชื่อมต่อ แอปจะส่งค่า Wi-Fi, ที่อยู่เซิร์ฟเวอร์ MQTT และโทเคนยืนยันตัวตนไปเก็บในหน่วยความจำถาวร (NVS) ของบอร์ด จากนั้นบอร์ดจะรีสตาร์ท ต่อ Wi-Fi ซิงก์เวลา NTP และขึ้นสถานะ ออนไลน์ ในระบบทันที")

add_callout_box(
    "การกดปุ่มบังคับเข้าโหมดตั้งค่าบลูทูธใหม่ (Reconfigure)",
    "หากต้องการเปลี่ยนรหัส Wi-Fi หรือย้ายกล่องไปห้องอื่นขณะที่บอร์ดไม่ได้เปิดบลูทูธอยู่ ให้เปิดฝากล่องแล้ว กดปุ่ม BOOT บนบอร์ด ESP32 ค้างไว้ แล้วกดปุ่ม RESET (หรือถอดเสียบไฟใหม่ขณะกด BOOT ค้างไว้ประมาณ 2 วินาที) บอร์ดจะเปิดสัญญาณบลูทูธให้ตั้งค่าใหม่ทันที",
    "warn"
)

add_sec_heading("7.2 ข้อกำหนดการติดตั้งหน้างานจริง")
add_clean_table(
    ["จุดติดตั้งหน้างาน", "วิธีติดตั้งที่ถูกต้อง"],
    [
        [
            "การคล้องห่วงวัดกระแส (CT Sensor)",
            "ง้างตัวล็อกห่วง CT สีดำออก แล้วนำไปคล้องรอบ สายไฟเส้นที่มีไฟ (Line: L) ของแอร์เพียงเส้นเดียว แล้วกดล็อกให้สนิท (ห้ามคล้องรวบทั้งสาย L และ N คู่กันเด็ดขาด เพราะกระแสไป-กลับจะหักล้างกันจนอ่านค่าวัตต์ได้ 0 W)"
        ],
        [
            "การวางหลอดยิงแสง IR LED\n(เฉพาะบอร์ดแอร์ IR)",
            "นำปลายสายหลอด IR LED ไปติดในตำแหน่งที่หันชี้เข้าหาช่องรับสัญญาณรีโมตของตัวแอร์โดยไม่มีสิ่งกีดขวาง"
        ],
        [
            "การติดกล่องเซ็นเซอร์ตรวจจับคน\n(เรดาร์ HLK-LD2410)",
            "ติดตั้งบนผนังห้องเรียนในมุมที่หันหน้ากล่องส่องไปยังบริเวณที่นั่งเรียนทั้งหมด หลีกเลี่ยงการหันหน้าเรดาร์เข้าหาพัดลมเพดานหรือม่านที่ปลิวลมแรง"
        ]
    ],
    col_widths_cm=[5.0, 12.0]
)

# ==============================================================================
# PAGE 12: ภาคผนวก ก ความหมายไฟแสดงสถานะ LED และหน้าจอ OLED
# ==============================================================================
doc.add_page_break()
add_chapter_heading("ก", "ภาคผนวก · ไฟแสดงสถานะและหน้าจอ OLED")
add_sec_heading("ก.1 ไฟแสดงสถานะ LED 3 สี บนบอร์ดแอร์ (ทั้งรุ่น IR และรุ่นสายแอร์)")
add_p("ด้านบนของกล่องบอร์ดแอร์มีหลอดไฟ LED 3 ดวง ช่วยให้ผู้ดูแลและช่างตรวจสอบการทำงานหน้าห้องได้ทันทีโดยไม่ต้องเปิดแอป")

add_clean_table(
    ["สีไฟ LED", "ขาพิน ESP32", "ความหมายของการแสดงผล"],
    [
        ["🟢 สีเขียว (สถานะแอร์)", "GPIO 25", "ติดค้าง = มิเตอร์ PZEM ตรวจพบว่าแอร์กำลังกินไฟทำงานอยู่ (กำลังไฟเกินเกณฑ์ 50 W ติดต่อกัน) · ดับ = แอร์ปิดอยู่"],
        ["🔴 สีแดง (นอกตารางเรียน)", "GPIO 26", "ติดค้าง = แอร์กำลังเปิดทำงานอยู่ในช่วงเวลาที่ไม่มีคาบเรียนหรือการจองห้อง · ดับ = ปกติ (อยู่ในคาบเรียน หรือแอร์ปิดอยู่)"],
        ["🟡 สีเหลือง (สถานะเครือข่าย)", "GPIO 27 (รุ่น IR)\nGPIO 23 (รุ่นสายแอร์)", "กะพริบทุก 400 ms = กำลังเชื่อมต่อ Wi-Fi หรืออยู่ในโหมดตั้งค่าบลูทูธ · ติดค้าง = เชื่อมต่อ Wi-Fi สำเร็จแล้ว"]
    ],
    col_widths_cm=[4.2, 3.3, 9.5]
)

add_sec_heading("ก.2 ไฟแสดงสถานะ LED 2 สี บนกล่องเซ็นเซอร์ตรวจจับคน (HLK-LD2410)")
add_clean_table(
    ["สีไฟ LED", "ขาพิน ESP32", "ความหมายของการแสดงผล"],
    [
        ["🟢 สีเขียว (สถานะระบบ)", "GPIO 25", "ติดค้าง = เชื่อมต่อ Wi-Fi และ MQTT ปกติ · กะพริบ = กำลังเชื่อมต่อ Wi-Fi/NTP หรือกำลังลองต่อใหม่ · ดับ = ยังไม่ได้ตั้งค่า หรืออยู่ในโหมดบลูทูธ"],
        ["🔴 สีแดง (ตรวจจับคน)", "GPIO 26", "ติดค้าง = ตรวจพบคนอยู่ในห้องเรียน (Occupied) · ดับ = ไม่มีคนในห้อง หรือเรดาร์ไม่ส่งข้อมูล"]
    ],
    col_widths_cm=[4.2, 3.3, 9.5]
)

add_sec_heading("ก.3 ข้อมูลที่แสดงบนหน้าจอ OLED 1.3 นิ้ว (ฝาหน้ากล่องบอร์ดแอร์)")
add_clean_table(
    ["ข้อความ / สัญลักษณ์บนจอ", "ความหมาย"],
    [
        ["ตัวเลขอุณหภูมิขนาดใหญ่ (เช่น 25°C)", "อุณหภูมิเป้าหมายที่สั่งตั้งค่าแอร์ไว้ล่าสุด (หรือแสดงคำว่า OFF เมื่อสั่งปิดแอร์)"],
        ["ค่ากำลังไฟและกระแส (W / V / A)", "ค่ากำลังไฟฟ้า (วัตต์), แรงดัน (โวลต์) และกระแส (แอมแปร์) ที่อ่านได้สดจากมิเตอร์ PZEM-004T"],
        ["ป้าย BUS และ COMP ON\n(เฉพาะบอร์ดสายแอร์ Daikin)", "BUS ทึบ = สายแอร์สื่อสารสองทางปกติ (ถ้าสายหลุดเกิน 15 วินาที จอจะสลับขึ้นหน้าเตือน) · COMP ON = แอร์รายงานว่าคอมเพรสเซอร์กำลังทำงานทำความเย็น"],
        ["โหมดพักหน้าจออัตโนมัติ (Screen Saver)", "เมื่อแอร์ปิดติดต่อกันเกิน 10 นาที หน้าจอจะดับพักเองอัตโนมัติเพื่อป้องกันพิกเซลจอ OLED เสื่อม (Burn-in) และจะติดขึ้นมาใหม่ทันทีที่มีการสั่งงาน"]
    ],
    col_widths_cm=[5.0, 12.0]
)

# ==============================================================================
# PAGE 13: ภาคผนวก ข การแก้ปัญหาฮาร์ดแวร์เบื้องต้น
# ==============================================================================
doc.add_page_break()
add_chapter_heading("ข", "ภาคผนวก · การแก้ปัญหาฮาร์ดแวร์เบื้องต้น")

add_clean_table(
    ["อาการที่พบ", "สาเหตุที่พบบ่อย", "วิธีตรวจสอบและแก้ไข"],
    [
        [
            "คอมไพล์โค้ดไม่ผ่าน แจ้งว่า Sketch too big (พื้นที่แฟลชไม่พอ)",
            "เลือก Partition Scheme ไม่ตรง หรือไฟล์ build_opt.h หายไปจากโฟลเดอร์",
            "ตรวจว่ามีไฟล์ build_opt.h อยู่ในโฟลเดอร์เดียวกับ .ino และตั้งค่า Tools → Partition Scheme เป็น No FS 4MB (2MB APP x2)"
        ],
        [
            "เสียบสาย USB แล้วมองไม่เห็น Port ใน Arduino IDE",
            "สาย USB Type-C ที่ใช้เป็นสายชาร์จอย่างเดียว (ไม่มีสายข้อมูล Data) หรือยังไม่ได้ลงไดรเวอร์ชิป USB",
            "เปลี่ยนเป็นสาย USB Type-C ที่รองรับการโอนข้อมูล และติดตั้งไดรเวอร์ CP210x หรือ CH340 ตามชิปบนบอร์ด ESP32"
        ],
        [
            "แอร์เปิดทำงานแล้ว แต่ค่ากำลังไฟในแอปขึ้น 0 W ตลอด",
            "คล้องห่วง CT รวบสายไฟทั้ง 2 เส้น (L+N) หรือต่อสาย TX/RX ของ PZEM สลับขา หรือไม่ได้ต่อไฟ 220V เข้าฝั่งซ้ายของ PZEM",
            "ตรวจว่าห่วง CT คล้องสาย L เพียงเส้นเดียว · ตรวจไฟ 220V ที่ขั้ว 1–2 ของ PZEM · ตรวจสาย UART (รุ่น IR ใช้ขา 16/18 ส่วนรุ่นสายแอร์ใช้ขา 18/14)"
        ],
        [
            "สั่งเปิดแอร์รุ่น IR จากแอปแล้ว แอร์ไม่ดังติ๊ด",
            "หลอด IR LED หันไม่ตรงช่องรับของแอร์ หรือต่อขาทรานซิสเตอร์ 2N2222 สลับด้าน",
            "ใช้กล้องมือถือส่องดูหลอด IR LED ขณะกดสั่งงาน (กล้องมือถือจะมองเห็นแสงสีม่วงกะพริบ) หากไม่กะพริบให้เช็คสายขา GPIO 17 และขา E/B/C ของ 2N2222"
        ],
        [
            "บอร์ดสายแอร์ Daikin หน้าจอขึ้นเตือนว่าสายไม่ตอบสนอง",
            "หันขั้วแถบดำของไดโอด 1N4148 ผิดด้าน หรือสลับสายสั่งงาน (ขาว) กับสายสถานะ (เหลือง) หรือลืมต่อ GND ร่วม",
            "เปิด Serial Monitor ที่ 115200 baud ตรวจค่า echo และเช็คว่าต่อสาย GND ของแอร์เข้ากับ GND ของ ESP32 แล้ว และแถบดำของ 1N4148 หันเข้าขา GPIO 33"
        ],
        [
            "เซ็นเซอร์ตรวจจับคนขึ้นว่า “เซ็นเซอร์ไม่พร้อมใช้งาน”",
            "สาย TX/RX ของเรดาร์ LD2410 หลวม หรือต่อไม่ไขว้กัน หรือจ่ายไฟเข้าขา 3.3V แทนที่จะเป็น 5V",
            "เรดาร์ LD2410 ต้องใช้ไฟเลี้ยงจากขั้ว 5V (VIN) เท่านั้น และตรวจว่าขา TX เรดาร์ต่อเข้า GPIO 32 ส่วนขา RX เรดาร์ต่อเข้า GPIO 33"
        ],
        [
            "หน้าจอ OLED 1.3 นิ้ว มืดสนิท แต่บอร์ดทำงานปกติ",
            "สาย SDA/SCL หลวม หรือสลับขากัน",
            "ตรวจสายจอ OLED ให้ตรงคือ VCC → 3V3, GND → GND, SDA → GPIO 21, SCL → GPIO 22 แล้วกดปุ่ม RESET บอร์ดใหม่"
        ]
    ],
    col_widths_cm=[4.4, 5.2, 7.4]
)

add_callout_box(
    "การตรวจสอบผ่าน Serial Monitor เมื่อต้องการหาสาเหตุเชิงลึก",
    "บอร์ดทั้ง 3 รูปแบบพิมพ์ข้อความสถานะการทำงานอย่างละเอียดออกทางพอร์ต USB ที่ความเร็ว 115200 baud เสมอ (เช่น สถานะการสแกนเจอจอ OLED, ค่าวัตต์ที่อ่านจาก PZEM, เฟรมสายแอร์ที่ส่ง-รับ และค่าพลังงานเรดาร์รายเกต Gate 0–8) หากพบปัญหาหน้างาน สามารถนำโน้ตบุ๊กไปเสียบสาย USB เพื่ออ่านข้อความวินิจฉัยได้ทันที",
    "info"
)

docx_ascii = os.path.join(out_dir, "RMUTT_Hardware_Assembly_Guide.docx")
docx_thai = os.path.join(out_dir, "คู่มือศึกษาโปรเจค", "คู่มือการประกอบและติดตั้งฮาร์ดแวร์_RMUTT_Save_Energy.docx")

doc.save(docx_ascii)
doc.save(docx_thai)
print("Generated style-matched Word (.docx) files successfully!")
